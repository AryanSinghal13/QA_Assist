import os
import io
import json
import streamlit as st
import pandas as pd
from google import genai
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ── API Key ───────────────────────────────────────────────────────────────────
API_KEY = os.getenv("GOOGLE_API_KEY", "")

# ── Shared Models & Config ────────────────────────────────────────────────────
@st.cache_resource
def configure_model():
    if not API_KEY:
        st.error("❌ No API key found. Add `GOOGLE_API_KEY` to your `.env` file.")
        st.stop()
    client = genai.Client(api_key=API_KEY)
    return client

def parse_response(text: str) -> list[dict]:
    text = text.strip()
    for fence in ("```json", "```"):
        if text.startswith(fence):
            text = text[len(fence):]
            break
    if text.endswith("```"):
        text = text[:-3]
    return json.loads(text.strip())

# ── Shared Prompts ────────────────────────────────────────────────────────────
BASE_PROMPT = """
You are a Senior QA Engineer with 10+ years of experience.
Generate comprehensive test cases covering:
1. Positive/Happy Path Scenarios
2. Negative/Error Path Scenarios
3. Edge Cases & Boundary Values
4. UI/UX and Accessibility considerations

Return ONLY a raw JSON array. DO NOT use markdown code fences. DO NOT include any text outside the array.

Example format:
[
  {
    "Module": "Login",
    "Test Case ID": "TC_LOG_001",
    "Scenario": "Valid credentials login",
    "Description": "1. Enter valid user\\n2. Enter valid pass\\n3. Click Login",
    "Expected Result": "User is redirected to dashboard.",
    "Priority": "High",
    "Type": "Positive"
  }
]

Make sure every JSON object contains exactly these keys:
  "Module"          – high-level feature area
  "Test Case ID"    – unique identifier (e.g. TC_001)
  "Scenario"        – brief title of the test
  "Description"     – numbered steps to reproduce
  "Expected Result" – what should happen
  "Priority"        – High | Medium | Low
  "Type"            – Positive | Negative | Edge Case | UI
"""

# ── Export & Display Helpers ──────────────────────────────────────────────────
def build_excel(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Test Cases")
        worksheet = writer.sheets["Test Cases"]
        
        # Define styles
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        wrap_alignment = Alignment(wrap_text=True, vertical="top")
        center_alignment = Alignment(horizontal="center", vertical="top")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                             top=Side(style='thin'), bottom=Side(style='thin'))

        # Style headers & set filter
        for col_num, cell in enumerate(worksheet[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = thin_border
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.freeze_panes = "A2"

        # Set column widths and alignments
        col_widths = {"A": 15, "B": 20, "C": 35, "D": 50, "E": 40, "F": 15, "G": 15}
        for col_letter, width in col_widths.items():
            worksheet.column_dimensions[col_letter].width = width

        # Apply styles to data rows
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
            for cell in row:
                cell.border = thin_border
                cell.alignment = wrap_alignment

            priority_idx = None
            for idx, col_name in enumerate(df.columns):
                if col_name == "Priority":
                    priority_idx = idx
                    break
            
            if priority_idx is not None:
                priority_cell = row[priority_idx]
                val = priority_cell.value
                if val == "High":
                    priority_cell.fill = PatternFill(start_color="FCA5A5", end_color="FCA5A5", fill_type="solid")
                    priority_cell.font = Font(color="7F1D1D", bold=True)
                elif val == "Medium":
                    priority_cell.fill = PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid")
                    priority_cell.font = Font(color="92400E", bold=True)
                elif val == "Low":
                    priority_cell.fill = PatternFill(start_color="A7F3D0", end_color="A7F3D0", fill_type="solid")
                    priority_cell.font = Font(color="065F46", bold=True)

    return buf.getvalue()


def show_results(df: pd.DataFrame, download_name: str = "test_cases.xlsx"):
    st.success(f"✅  {len(df)} test cases generated!")

    # Filters row
    col_p, col_t, col_m = st.columns(3)
    with col_p:
        prios   = ["All"] + sorted(df["Priority"].dropna().unique().tolist())
        sel_p   = st.selectbox("Filter Priority", prios, key=f"p_{download_name}")
    with col_t:
        types   = ["All"] + sorted(df["Type"].dropna().unique().tolist())
        sel_t   = st.selectbox("Filter Type",     types, key=f"t_{download_name}")
    with col_m:
        modules = ["All"] + sorted(df["Module"].dropna().unique().tolist())
        sel_m   = st.selectbox("Filter Module",   modules, key=f"m_{download_name}")

    filtered = df.copy()
    if sel_p != "All": filtered = filtered[filtered["Priority"] == sel_p]
    if sel_t != "All": filtered = filtered[filtered["Type"]     == sel_t]
    if sel_m != "All": filtered = filtered[filtered["Module"]   == sel_m]

    # Summary metrics
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Total", len(filtered))
    m2.metric("🔴 High",   len(filtered[filtered["Priority"] == "High"]))
    m3.metric("🟡 Medium", len(filtered[filtered["Priority"] == "Medium"]))
    m4.metric("🟢 Low",    len(filtered[filtered["Priority"] == "Low"]))

    def priority_color(val):
        return {
            "High":   "background-color:#7f1d1d;color:#fca5a5",
            "Medium": "background-color:#78350f;color:#fcd34d",
            "Low":    "background-color:#14532d;color:#86efac",
        }.get(val, "")

    styled = filtered.style.applymap(priority_color, subset=["Priority"])
    st.dataframe(styled, use_container_width=True, height=420)

    dl1, dl2 = st.columns(2)
    with dl1:
        st.download_button("📥 Download Excel", build_excel(filtered),
                           file_name=download_name,
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with dl2:
        st.download_button("📄 Download CSV", filtered.to_csv(index=False, encoding='utf-8-sig').encode('utf-8-sig'),
                           file_name=download_name.replace(".xlsx", ".csv"), mime="text/csv")
